"""
Extractor de las fuentes históricas de Gabriel (F01–F10) a JSON "crudo".

Lee los originales (xlsx/xls/pdf) que viven FUERA de Git y escribe, por archivo, un JSON con todas las
filas físicas que tienen contenido, sin interpretar nada: la interpretación (layouts, DNI/CUIL, eventos)
es del importador TypeScript (lib/imports/gabriel). Así el importador no depende de librerías de Excel/PDF.

Los JSON de salida contienen datos personales: van SIEMPRE fuera de Git (carpeta /data/ está ignorada).

Uso:
  python tools/gabriel-extract.py <carpeta_raw> <carpeta_salida>
  ej.: python tools/gabriel-extract.py C:/Users/usuario/Documents/sutecba-fuentes/raw data/gabriel/extracted

Dependencias: openpyxl, pandas, xlrd (para .xls), pymupdf (fitz, para los PDF).
Formato de celda: JSON escalar; una fecha se emite como {"$date": "AAAA-MM-DD"} o {"$datetime": "ISO"}.
"""
import datetime as dt
import hashlib
import json
import os
import sys
import warnings

warnings.filterwarnings("ignore")

# Código estable de cada fuente (docs/importacion-gabriel-mapeo.md). La coincidencia es por nombre de archivo.
FILES = [
    ("F01", "R.C.P -Cruz Malta (Respuestas).pdf"),
    ("F02", "52010-RCP CRUZ MALTA.xlsx"),
    ("F03", "AGC-CAPACITACION 2026 (Respuestas).pdf"),
    ("F04", "A.G.C - Primeros auxilio psicologicos (Respuestas).xlsx"),
    ("F05", "52469-INTELIGENCIA EMOCIONAL EN LA ORGANIZACION - A.G.C.xlsx"),
    ("F06", "Padron Gral x Repart.xlsx"),
    ("F07", "Padrón PG.xls"),
    ("F08", "OFTALMO 2026.xlsx"),
    ("F09", "Oftalmo Teatro Colón (Respuestas).xlsx"),
    ("F10", "abogados unificado.xlsx"),
]


def cell_value(v):
    if v is None:
        return None
    if isinstance(v, float):
        if v != v:  # NaN
            return None
        return int(v) if v.is_integer() else v
    if isinstance(v, dt.datetime):
        if v.hour == 0 and v.minute == 0 and v.second == 0:
            return {"$date": v.date().isoformat()}
        return {"$datetime": v.isoformat()}
    if isinstance(v, dt.date):
        return {"$date": v.isoformat()}
    if hasattr(v, "to_pydatetime"):  # pandas.Timestamp
        return cell_value(v.to_pydatetime())
    if isinstance(v, (int, bool)):
        return v
    s = str(v)
    return s if s.strip() != "" else None


def trim(cells):
    out = list(cells)
    while out and out[-1] is None:
        out.pop()
    return out


def sheet_rows(rows_iter):
    rows = []
    for i, row in enumerate(rows_iter, start=1):
        cells = trim([cell_value(c) for c in row])
        if any(c is not None for c in cells):
            rows.append({"n": i, "cells": cells})
    return rows


def read_xlsx(path):
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    return [{"name": ws.title, "rows": sheet_rows(ws.iter_rows(values_only=True))} for ws in wb.worksheets]


def read_xls(path):
    import pandas as pd

    book = pd.ExcelFile(path, engine="xlrd")
    sheets = []
    for name in book.sheet_names:
        df = book.parse(name, header=None)
        sheets.append({"name": name, "rows": sheet_rows(df.itertuples(index=False, name=None))})
    return sheets


def read_pdf(path):
    import fitz

    doc = fitz.open(path)
    sheets = []
    for pno, page in enumerate(doc, start=1):
        for tno, table in enumerate(page.find_tables().tables, start=1):
            data = table.extract()
            rows = sheet_rows(data)
            sheets.append({"name": f"p{pno}" + (f"-t{tno}" if tno > 1 else ""), "rows": rows})
    return sheets


def main(raw_dir, out_dir):
    os.makedirs(out_dir, exist_ok=True)
    for code, name in FILES:
        path = os.path.join(raw_dir, name)
        if not os.path.exists(path):
            print(f"[extract] {code}: NO ENCONTRADO ({name})")
            continue
        with open(path, "rb") as fh:
            blob = fh.read()
        ext = os.path.splitext(name)[1].lower()
        sheets = read_pdf(path) if ext == ".pdf" else read_xls(path) if ext == ".xls" else read_xlsx(path)
        payload = {
            "fileCode": code,
            "fileName": name,
            "sha256": hashlib.sha256(blob).hexdigest(),
            "sizeBytes": len(blob),
            "sheets": sheets,
        }
        target = os.path.join(out_dir, f"{code}.json")
        with open(target, "w", encoding="utf-8") as out:
            json.dump(payload, out, ensure_ascii=False)
        total = sum(len(s["rows"]) for s in sheets)
        print(f"[extract] {code}: {len(sheets)} hoja(s), {total} filas con contenido -> {target}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(1)
    main(sys.argv[1], sys.argv[2])
