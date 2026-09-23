"""
Escribe un libro Excel de revisión humana a partir de una especificación JSON.

    python tools/review-xlsx.py <spec.json> <salida.xlsx>

spec = { "sheets": [ { "name", "note"?, "columns": [{"key","header","width"?,"wrap"?}], "rows": [ {key: valor} ],
                       "validations": [{"key": "decision", "options": ["A","B"]}], "freeze"?: "C2", "editable": ["decision", ...] } ] }

Las columnas «editables» (las de decisión) se resaltan; el resto es evidencia. Los datos personales quedan SOLO en el
.xlsx (fuera de Git); este script no imprime ninguna fila.
"""
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
EDIT_FILL = PatternFill("solid", fgColor="FFF2CC")
EDIT_HEADER_FILL = PatternFill("solid", fgColor="B45F06")


def main(spec_path, out_path):
    spec = json.load(open(spec_path, encoding="utf-8"))
    wb = Workbook()
    wb.remove(wb.active)
    for sheet in spec["sheets"]:
        ws = wb.create_sheet(sheet["name"][:31])
        cols = sheet["columns"]
        editable = set(sheet.get("editable", []))
        first_row = 1
        if sheet.get("note"):
            ws.cell(row=1, column=1, value=sheet["note"]).font = Font(italic=True, color="7F0000")
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(4, min(len(cols), 8)))
            ws.cell(row=1, column=1).alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[1].height = 48
            first_row = 2
        for i, col in enumerate(cols, start=1):
            cell = ws.cell(row=first_row, column=i, value=col["header"])
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = EDIT_HEADER_FILL if col["key"] in editable else HEADER_FILL
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(i)].width = col.get("width", 18)
        for r, row in enumerate(sheet["rows"], start=first_row + 1):
            for i, col in enumerate(cols, start=1):
                value = row.get(col["key"])
                if isinstance(value, list):
                    value = "\n".join(str(v) for v in value)
                cell = ws.cell(row=r, column=i, value=value)
                cell.alignment = Alignment(wrap_text=col.get("wrap", True), vertical="top")
                if col["key"] in editable:
                    cell.fill = EDIT_FILL
        last = first_row + len(sheet["rows"])
        if sheet["rows"]:
            ws.auto_filter.ref = f"A{first_row}:{get_column_letter(len(cols))}{last}"
        ws.freeze_panes = sheet.get("freeze", f"A{first_row + 1}")
        for v in sheet.get("validations", []):
            idx = [c["key"] for c in cols].index(v["key"]) + 1
            dv = DataValidation(type="list", formula1='"' + ",".join(v["options"]) + '"', allow_blank=True, showErrorMessage=True)
            ws.add_data_validation(dv)
            letter = get_column_letter(idx)
            dv.add(f"{letter}{first_row + 1}:{letter}{max(last, first_row + 1) + 200}")
    wb.save(out_path)
    print({"hojas": [s["name"] for s in spec["sheets"]], "filas": [len(s["rows"]) for s in spec["sheets"]]})


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
