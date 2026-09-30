#!/usr/bin/env python
"""Extractor de las 22 fuentes de las tandas 1 y 2 -> JSON por archivo (solo lectura de los originales).

Uso:  python tools/tandas-extract.py --manifest <manifest.json> --out data/tandas/extracted

El manifest es una lista [{"code": "T01", "path": "...", "fileName": "..."}]. Cada JSON tiene la misma forma que los
extractos históricos (ver lib/imports/gabriel/types.ts: ExtractedFile), con fileCode "T01".."T22":
  {fileCode, fileName, sha256, sizeBytes, sheets:[{name, rows:[{n, cells}]}]}
`n` es el número de fila FÍSICA (1-based) de la hoja. No se descarta ninguna fila (ni vacía): la conciliación necesita
contar filas físicas. Fechas -> {"$date": "AAAA-MM-DD"} / {"$datetime": "AAAA-MM-DDTHH:MM:SS"}; números enteros -> int.
Los números no enteros se conservan como float (la pérdida de precisión la detecta el normalizador en TypeScript).
"""
import argparse
import datetime as dt
import hashlib
import json
import os
import sys


def cell(v):
    if v is None:
        return None
    if isinstance(v, bool):
        return v
    if isinstance(v, dt.datetime):
        if v.hour == 0 and v.minute == 0 and v.second == 0 and v.microsecond == 0:
            return {"$date": v.strftime("%Y-%m-%d")}
        return {"$datetime": v.strftime("%Y-%m-%dT%H:%M:%S")}
    if isinstance(v, dt.date):
        return {"$date": v.strftime("%Y-%m-%d")}
    if isinstance(v, float):
        if v != v:
            return None
        if v == int(v) and abs(v) < 1e15:
            return int(v)
        return v
    if isinstance(v, str):
        s = v.replace(" ", " ")
        return s if s.strip() != "" else None
    return v


def read_xlsx(path):
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        rows = []
        for i, r in enumerate(ws.iter_rows(values_only=True), 1):
            rows.append({"n": i, "cells": [cell(c) for c in r]})
        yield ws.title, rows


def read_xls(path):
    import xlrd

    wb = xlrd.open_workbook(path)
    for sh in wb.sheets():
        rows = []
        for i in range(sh.nrows):
            out = []
            for j in range(sh.ncols):
                c = sh.cell(i, j)
                if c.ctype == xlrd.XL_CELL_DATE:
                    out.append(cell(xlrd.xldate.xldate_as_datetime(c.value, wb.datemode)))
                elif c.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    out.append(None)
                elif c.ctype == xlrd.XL_CELL_NUMBER:
                    out.append(cell(float(c.value)))
                else:
                    out.append(cell(c.value))
            rows.append({"n": i + 1, "cells": out})
        yield sh.name, rows


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--manifest", required=True)
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    os.makedirs(a.out, exist_ok=True)
    for item in json.load(open(a.manifest, encoding="utf-8")):
        path = item["path"]
        data = open(path, "rb").read()
        reader = read_xls if path.lower().endswith(".xls") else read_xlsx
        sheets = [{"name": n, "rows": r} for n, r in reader(path)]
        doc = {
            "fileCode": item["code"],
            "fileName": item["fileName"],
            "sha256": hashlib.sha256(data).hexdigest(),
            "sizeBytes": len(data),
            "sheets": sheets,
        }
        with open(os.path.join(a.out, item["code"] + ".json"), "w", encoding="utf-8") as fh:
            json.dump(doc, fh, ensure_ascii=False)
        print(item["code"], item["fileName"][:50], sum(len(s["rows"]) for s in sheets), "filas", file=sys.stderr)


if __name__ == "__main__":
    main()
