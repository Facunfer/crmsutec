"""
Extractor del catálogo organizacional (SUTECBA_Estructura_Oficial_Alias_Gabriel_vN.xlsx) a JSON.

    python tools/org-catalog-extract.py <archivo.xlsx> <salida.json>

Lee las hojas Organismos_Oficiales, Alias_Reparticion, Alias_Area_Interna, Pendientes y Decisiones_V2 y deja
todas las celdas como texto. Guarda el SHA-256 del .xlsx para poder verificar después que el catálogo que se
carga es exactamente el que se revisó. El JSON queda fuera de Git (/data/ está ignorado).
"""
import hashlib
import json
import os
import sys

import openpyxl

SHEETS = ["Organismos_Oficiales", "Alias_Reparticion", "Alias_Area_Interna", "Pendientes", "Decisiones_V2"]


def main(src, out):
    blob = open(src, "rb").read()
    wb = openpyxl.load_workbook(src, data_only=True)
    sheets = {}
    for name in SHEETS:
        ws = wb[name]
        rows = list(ws.iter_rows(values_only=True))
        header = [str(h) for h in rows[0]]
        sheets[name] = [
            {h: (None if c is None else str(c)) for h, c in zip(header, r)}
            for r in rows[1:]
            if any(c is not None for c in r)
        ]
    readme = {str(r[0]): r[1] for r in wb["README"].iter_rows(values_only=True) if r[0]} if "README" in wb.sheetnames else {}
    payload = {
        "fileName": os.path.basename(src),
        "sha256": hashlib.sha256(blob).hexdigest(),
        "version": readme.get("Versión"),
        "sheets": sheets,
    }
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    with open(out, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False)
    print({k: len(v) for k, v in sheets.items()}, payload["sha256"][:12], payload["version"])


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
